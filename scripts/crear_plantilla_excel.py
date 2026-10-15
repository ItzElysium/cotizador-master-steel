"""Crea data/plantilla_lista_precios.xlsx: la plantilla que llena el cliente con su lista de precios.

    python scripts/crear_plantilla_excel.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from cotizador.catalogo import COLUMNAS, leer_archivo  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
ARIAL = "Arial"
OSCURO, ACENTO, AMARILLO = "1F2937", "C2410C", "FFF7CC"

AYUDA = {
    "sku": "Código único, sin / \\ # ?. Ej: FC-0500",
    "categoria": "Familia: Fierro corrugado, Plancha LAC, Tubo LAC cuadrado...",
    "descripcion": "Como debe salir en la cotización",
    "sinonimos": "Cómo lo piden los clientes, separados por |  Ej: fierro de media|varilla 1/2",
    "unidad": "Cómo se vende: varilla 9 m, plancha 1.20 x 2.40 m, tubo 6 m, kg",
    "peso_kg": "Peso teórico de UNA unidad de venta (1 si se vende por kg)",
    "precio_usd": "Precio en US$ SIN IGV, por unidad o por kg según base_precio",
    "base_precio": "unidad = precio por varilla/plancha/tubo · kg = precio por kilo",
    "stock": "disponible · consultar · agotado",
    "actualizado": "Fecha del precio (AAAA-MM-DD). Si pasa de 7 días, el vendedor recibe alerta",
}
EDITABLES_FRECUENTES = {"precio_usd", "stock", "actualizado"}
ANCHOS = {"sku": 12, "categoria": 20, "descripcion": 52, "sinonimos": 48, "unidad": 20, "peso_kg": 10,
          "precio_usd": 12, "base_precio": 12, "stock": 13, "actualizado": 13}


def main() -> None:
    productos, _ = leer_archivo(RAIZ / "data" / "catalogo_ejemplo.csv")
    libro = Workbook()

    hoja = libro.active
    hoja.title = "lista_precios"
    borde = Border(bottom=Side(style="thin", color="D1D5DB"))
    for col, nombre in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=col, value=nombre)
        celda.font = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=ACENTO if nombre in EDITABLES_FRECUENTES else OSCURO)
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.comment = Comment(AYUDA[nombre], "DataPrix")
        hoja.column_dimensions[celda.column_letter].width = ANCHOS[nombre]
    hoja.row_dimensions[1].height = 22

    for fila, p in enumerate(productos, start=2):
        datos = p.a_dict()
        for col, nombre in enumerate(COLUMNAS, start=1):
            valor = datos[nombre]
            if nombre in {"peso_kg", "precio_usd"}:
                valor = float(valor)
            celda = hoja.cell(row=fila, column=col, value=valor)
            celda.font = Font(name=ARIAL, size=10, color="0000FF" if nombre in EDITABLES_FRECUENTES else "000000")
            celda.border = borde
            if nombre in EDITABLES_FRECUENTES:
                celda.fill = PatternFill("solid", fgColor=AMARILLO)
            if nombre == "precio_usd":
                celda.number_format = '"US$" #,##0.00'
            elif nombre == "peso_kg":
                celda.number_format = "#,##0.00"
    ultima = len(productos) + 500
    for nombre, opciones in {"base_precio": "unidad,kg", "stock": "disponible,consultar,agotado"}.items():
        letra = hoja.cell(row=1, column=COLUMNAS.index(nombre) + 1).column_letter
        validacion = DataValidation(type="list", formula1=f'"{opciones}"', allow_blank=False,
                                    error="Elige un valor de la lista", errorTitle="Valor no válido")
        hoja.add_data_validation(validacion)
        validacion.add(f"{letra}2:{letra}{ultima}")
    hoja.freeze_panes = "B2"
    hoja.auto_filter.ref = f"A1:{hoja.cell(row=1, column=len(COLUMNAS)).column_letter}{len(productos) + 1}"

    ayuda = libro.create_sheet("instrucciones")
    ayuda.column_dimensions["A"].width = 16
    ayuda.column_dimensions["B"].width = 95
    lineas = [
        ("Lista de precios del cotizador IA", ""),
        ("", ""),
        ("Qué editar", "Las columnas naranjas/amarillas (precio_usd, stock, actualizado) cambian seguido. "
                       "Las demás se llenan una vez por producto."),
        ("Ejemplo", "Las filas de la hoja lista_precios son un EJEMPLO con pesos teóricos y precios ficticios: "
                    "reemplázalas por tus productos y precios reales."),
        ("Precios", "En dólares y SIN IGV. El sistema agrega IGV, descuentos y el equivalente en soles."),
        ("Sinónimos", "Es lo que hace que el asistente entienda la jerga. Ej: fierro de media|varilla 1/2|corrugado 1/2"),
        ("Pesos", "Peso teórico por unidad de venta. Sirve para convertir pedidos en kilos/toneladas a varillas."),
        ("Subir", "python scripts/cargar_catalogo.py lista.xlsx --url https://<app>.azurewebsites.net --clave <clave>"),
        ("Cambio rápido", "Por WhatsApp (número de administrador): /precio FC-0500 8.90  ·  /stock FC-0500 agotado"),
    ]
    for i, (a, b) in enumerate(lineas, start=1):
        ayuda.cell(row=i, column=1, value=a).font = Font(name=ARIAL, bold=True, size=14 if i == 1 else 10,
                                                          color=OSCURO)
        celda = ayuda.cell(row=i, column=2, value=b)
        celda.font = Font(name=ARIAL, size=10)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    destino = RAIZ / "data" / "plantilla_lista_precios.xlsx"
    libro.save(destino)
    print(f"Plantilla creada: {destino} ({len(productos)} productos de ejemplo)")


if __name__ == "__main__":
    main()
