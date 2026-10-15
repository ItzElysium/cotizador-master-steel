"""Catálogo de productos: carga, validación y búsqueda tolerante a jerga."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .modelos import Producto
from .texto import es_numerico, normalizar, tokens

COLUMNAS = ["sku", "categoria", "descripcion", "sinonimos", "unidad", "peso_kg",
            "precio_usd", "base_precio", "stock", "actualizado"]
_CARACTERES_PROHIBIDOS_SKU = set('/\\#?')


@dataclass
class Coincidencia:
    producto: Producto
    puntaje: float


class ErrorCatalogo(ValueError):
    pass


def leer_csv(texto: str, existentes: dict[str, Producto] | None = None,
             hoy: str = "") -> tuple[list[Producto], list[str]]:
    """Lee el CSV de la lista de precios. Devuelve (productos válidos, errores).

    Si se pasan `existentes`, se admite una carga parcial (ej. solo `sku,precio_usd`):
    las columnas que no vienen se toman del producto actual.
    """
    texto = texto.lstrip("﻿")
    muestra = texto[:2000]
    delimitador = ";" if muestra.count(";") > muestra.count(",") else ","
    lector = csv.DictReader(io.StringIO(texto), delimiter=delimitador)
    lector.fieldnames = [c.strip().lower() for c in (lector.fieldnames or [])]
    obligatorias = ("sku", "precio_usd") if existentes is not None else ("sku", "descripcion", "peso_kg", "precio_usd")
    faltantes = [c for c in obligatorias if c not in lector.fieldnames]
    if faltantes:
        raise ErrorCatalogo(f"Faltan columnas obligatorias: {', '.join(faltantes)}")
    ausentes = [c for c in COLUMNAS if c not in lector.fieldnames]
    productos: list[Producto] = []
    errores: list[str] = []
    vistos: set[str] = set()
    for n, fila in enumerate(lector, start=2):
        fila = {k.strip(): (v or "").strip() for k, v in fila.items() if k}
        sku = fila.get("sku", "").upper()
        if not sku:
            continue
        fila["sku"] = sku
        if any(c in _CARACTERES_PROHIBIDOS_SKU for c in sku):
            errores.append(f"Fila {n}: el SKU '{sku}' no puede contener / \\ # ?")
            continue
        if sku in vistos:
            errores.append(f"Fila {n}: SKU duplicado '{sku}'")
            continue
        previo = (existentes or {}).get(sku)
        if previo and ausentes:
            base = previo.a_dict()
            for campo in ausentes:
                fila[campo] = base[campo]
            if "actualizado" in ausentes and hoy and fila.get("precio_usd") not in (None, "", str(previo.precio_usd)):
                fila["actualizado"] = hoy
        if not fila.get("descripcion"):
            errores.append(f"Fila {n} ({sku}): producto nuevo sin descripción")
            continue
        try:
            for campo in ("peso_kg", "precio_usd"):
                fila[campo] = str(Decimal(fila.get(campo, "0").replace(",", ".") or "0"))
            producto = Producto.de_dict(fila)
        except (InvalidOperation, KeyError) as exc:
            errores.append(f"Fila {n} ({sku}): valor inválido ({exc})")
            continue
        if producto.precio_usd <= 0:
            errores.append(f"Fila {n} ({sku}): precio_usd debe ser mayor a 0")
            continue
        if producto.peso_kg <= 0:
            errores.append(f"Fila {n} ({sku}): peso_kg debe ser mayor a 0")
            continue
        if producto.base_precio not in {"unidad", "kg"}:
            errores.append(f"Fila {n} ({sku}): base_precio debe ser 'unidad' o 'kg'")
            continue
        if producto.stock not in {"disponible", "consultar", "agotado"}:
            producto.stock = "consultar"
        vistos.add(sku)
        productos.append(producto)
    return productos, errores


def leer_archivo(ruta: str | Path) -> tuple[list[Producto], list[str]]:
    ruta = Path(ruta)
    if ruta.suffix.lower() in {".xlsx", ".xlsm"}:
        return leer_csv(excel_a_csv(ruta))
    return leer_csv(ruta.read_text(encoding="utf-8"))


def excel_a_csv(ruta: Path) -> str:
    from openpyxl import load_workbook  # dependencia solo para scripts locales

    libro = load_workbook(ruta, data_only=True, read_only=True)
    hoja = libro["lista_precios"] if "lista_precios" in libro.sheetnames else libro.worksheets[0]
    salida = io.StringIO()
    escritor = csv.writer(salida)
    for fila in hoja.iter_rows(values_only=True):
        if fila is None or all(v is None for v in fila):
            continue
        escritor.writerow(["" if v is None else v for v in fila])
    return salida.getvalue()


def a_csv(productos: list[Producto]) -> str:
    salida = io.StringIO()
    escritor = csv.DictWriter(salida, fieldnames=COLUMNAS)
    escritor.writeheader()
    for p in productos:
        escritor.writerow(p.a_dict())
    return salida.getvalue()


class Catalogo:
    """Índice en memoria del catálogo con búsqueda por tokens normalizados."""

    def __init__(self, productos: list[Producto]):
        self.productos = {p.sku.upper(): p for p in productos}
        self._indice: dict[str, tuple[set[str], set[str], str]] = {}
        for p in productos:
            textos = [p.descripcion, p.categoria, *p.sinonimos]
            toks = set()
            for t in textos:
                toks.update(tokens(t))
            numericos = {t for t in toks if es_numerico(t)}
            self._indice[p.sku.upper()] = (toks, numericos, normalizar(" | ".join(textos)))

    def __len__(self) -> int:
        return len(self.productos)

    def obtener(self, sku: str) -> Producto | None:
        return self.productos.get((sku or "").strip().upper())

    def buscar(self, consulta: str, limite: int = 6) -> list[Coincidencia]:
        directo = self.obtener(consulta)
        if directo:
            return [Coincidencia(directo, 100.0)]
        q_toks = tokens(consulta)
        if not q_toks:
            return []
        q_num = [t for t in q_toks if es_numerico(t)]
        q_pal = [t for t in q_toks if not es_numerico(t)]
        q_norm = normalizar(consulta)
        resultados: list[Coincidencia] = []
        for sku, (toks, numericos, texto) in self._indice.items():
            puntaje = 0.0
            for t in q_pal:
                if t in toks:
                    puntaje += 1.0
                elif len(t) >= 4 and any(x.startswith(t[:4]) for x in toks):
                    puntaje += 0.5
            for t in q_num:
                if t in numericos:
                    puntaje += 3.0
                else:
                    puntaje -= 2.0  # una medida que no coincide es mala señal
            if q_norm and q_norm in texto:
                puntaje += 2.0
            if puntaje > 0:
                resultados.append(Coincidencia(self.productos[sku], puntaje))
        resultados.sort(key=lambda c: (-c.puntaje, c.producto.sku))
        return resultados[:limite]

    def resumen_para_prompt(self) -> str:
        """Tabla compacta (sin precios) para que el agente conozca el catálogo."""
        lineas = ["sku | descripción | se vende por | también le dicen"]
        for p in self.productos.values():
            sinonimos = ", ".join(p.sinonimos[:3])
            lineas.append(f"{p.sku} | {p.descripcion} | {p.unidad} | {sinonimos}")
        return "\n".join(lineas)
